"""Generator szablonu XLSX dla aptek (stany i ceny surowców farmaceutycznych konopi).

Użycie (z katalogu głównego repo):
    python tools/build_template.py strains.json Szablon_stany_i_ceny_konopie.xlsx

Wymaga: openpyxl (pip install openpyxl).
Po zmianie listy odmian w strains.json wygeneruj szablon ponownie i podnieś
template_version (arkusz "meta"), żeby automat rozpoznał nieaktualne pliki od aptek.

Struktura arkusza "Stany i ceny" (dla automatu odczytującego pliki od aptek):
    wiersze danych od wiersza 7 (meta.data_first_row)
    A ID (klucz: producent_nazwa_thc_gramaturag) | B producent | C nazwa | D THC | E opakowanie (g)
    F cena (zł/op.) | G stan z listy | H liczba opakowań (opcjonalnie, ma pierwszeństwo przed G) | I nowość (TAK)
    J, K ukryte kolumny pomocnicze (numer do wydruku, stan efektywny)
Stany: "dużo (15+)", "średnio (5–14)", "mało (1–4)", "brak". Pusta cena = apteka nie prowadzi.
"""
import json, re, sys, unicodedata
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment

src, out = sys.argv[1], sys.argv[2]
data = json.load(open(src, encoding="utf-8"))
strains = data["strains"]
SRC_DATE = data["meta"].get("last_updated", "")

EXCLUDE = {("Aurora", "Delahaze")}  # wycofana XI 2025, brak w obrocie


def slug(s):
    s = unicodedata.normalize("NFKD", s.replace("ł", "l").replace("Ł", "L"))
    s = s.encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


rows = []
for x in strains:
    if (x["producer"], x["name"]) in EXCLUDE:
        continue
    packs = sorted({p.get("package_g") for p in x.get("prices", []) if p.get("package_g")})
    raw = x["thc"].replace("%", "").strip()
    # zakres THC (np. 27–30%) = osobne produkty w aptece -> osobny wiersz na każdy wariant
    thcs = [t.strip() for t in re.split(r"[–-]", raw)] if re.search(r"[–-]", raw) else [raw]
    for thc in thcs:
        for pk in (packs or [None]):
            rid = f"{slug(x['producer'])}_{slug(x['name'])}_{slug(thc)}_{pk if pk else 'x'}g"
            t = (x.get("type") or "").strip().lower()
            typ = "INDICA" if t.startswith("indica") else "SATIVA" if t.startswith("sativa") else "HYBRYDA" if t else ""
            rows.append(dict(id=rid, producer=x["producer"], name=x["name"], thc=f"{thc}%", pack=pk,
                             split=len(thcs) > 1, typ=typ))
rows.sort(key=lambda r: (r["producer"].lower(), r["name"].lower(), r["thc"], r["pack"] or 0))
assert len({r["id"] for r in rows}) == len(rows), "duplicate ids"

F = "Arial"
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
fill_in = PatternFill("solid", fgColor="FFF2CC")      # pola do wypełnienia
fill_head = PatternFill("solid", fgColor="1F4E78")
fill_lock = PatternFill("solid", fgColor="F2F2F2")
fill_group = PatternFill("solid", fgColor="DDEBF7")
unlocked = Protection(locked=False)

wb = Workbook()

# ---------------- Arkusz 1: Stany i ceny ----------------
ws = wb.active
ws.title = "Stany i ceny"
ws.sheet_view.showGridLines = False

ws["A1"] = "Stany i ceny surowców farmaceutycznych konopi"
ws["A1"].font = Font(name=F, size=14, bold=True, color="1F4E78")
ws["A2"] = "Wypełnij tylko żółte pola. Pusta cena = apteka nie prowadzi tej odmiany."
ws["A2"].font = Font(name=F, size=10, italic=True, color="595959")

ws["B3"] = "Apteka:"
ws["B4"] = "Data stanu:"
for c in ("B3", "B4"):
    ws[c].font = Font(name=F, size=11, bold=True)
    ws[c].alignment = Alignment(horizontal="right")
for c in ("C3", "C4"):
    ws[c].fill = fill_in
    ws[c].border = border
    ws[c].protection = unlocked
    ws[c].font = Font(name=F, size=11)
ws.merge_cells("C3:E3")
ws.merge_cells("C4:E4")
ws["C4"].number_format = "DD.MM.YYYY"
ws["C3"].comment = Comment("Pełna nazwa i adres apteki, np. Apteka Niezapominajka, ul. Rejtana, Rzeszów", "lek. Michał Guzik")
ws["C4"].comment = Comment("Dzień, z którego pochodzą stany, np. 20.10.2026", "lek. Michał Guzik")

HDR = 6
headers = ["ID (nie zmieniać)", "Producent", "Nazwa odmiany", "THC", "Opakowanie (g)",
           "Cena (zł/op.)", "Stan", "Liczba op. (opcjonalnie)", "Nowość"]
widths = [30, 22, 32, 9, 14, 14, 16, 14, 10]
for i, (h, w) in enumerate(zip(headers, widths), start=1):
    c = ws.cell(row=HDR, column=i, value=h)
    c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
    c.fill = fill_head
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border
    ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[HDR].height = 30

r0 = HDR + 1
prev = None
for i, r in enumerate(rows):
    n = r0 + i
    vals = [r["id"], r["producer"], r["name"], r["thc"], r["pack"]]
    for col, v in enumerate(vals, start=1):
        c = ws.cell(row=n, column=col, value=v)
        c.font = Font(name=F, size=10, color="808080" if col == 1 else "000000",
                      bold=(col == 3))
        c.border = border
        c.fill = fill_lock
        c.alignment = Alignment(horizontal="center" if col in (4, 5) else "left", vertical="center")
    if r.get("split"):
        ws.cell(row=n, column=4).comment = Comment(
            "Odmiana występuje w kilku stężeniach THC. Każde stężenie ma osobny wiersz z własną ceną. "
            "Inne stężenie niż na liście wpisz w arkuszu „Inna odmiana”.", "lek. Michał Guzik")
    if r["pack"] is None:
        c = ws.cell(row=n, column=5)
        c.fill = fill_in
        c.protection = unlocked
        c.comment = Comment("Gramatura nieznana w bazie: wpisz wielkość opakowania (5, 10 lub 15 g)", "lek. Michał Guzik")
    for col in (6, 7):
        c = ws.cell(row=n, column=col)
        c.fill = fill_in
        c.border = border
        c.protection = unlocked
        c.font = Font(name=F, size=10, color="0000FF")
        c.alignment = Alignment(horizontal="center")
    ws.cell(row=n, column=6).number_format = '#,##0.00 "zł"'
    qn = ws.cell(row=n, column=8)
    qn.fill = PatternFill("solid", fgColor="FFFBEA"); qn.border = border; qn.protection = unlocked
    qn.font = Font(name=F, size=10, color="0000FF"); qn.alignment = Alignment(horizontal="center")
    qn.number_format = "0"
    # ukryty stan efektywny: z liczby opakowań (jeśli wpisana, ma pierwszeństwo), inaczej z listy
    ws.cell(row=n, column=11, value=(f'=IF(F{n}="","",IF(ISNUMBER(H{n}),'
        f'IF(H{n}=0,"brak",IF(H{n}<=4,"mało (1–4)",IF(H{n}<=14,"średnio (5–14)","dużo (15+)"))),G{n}))')).font = Font(name=F, size=8, color="BFBFBF")
    nw = ws.cell(row=n, column=9)
    nw.fill = fill_in; nw.border = border; nw.protection = unlocked
    nw.font = Font(name=F, size=10, color="0000FF", bold=True); nw.alignment = Alignment(horizontal="center")
last = r0 + len(rows) - 1

# walidacja
dv_price = DataValidation(type="decimal", operator="between", formula1="1", formula2="5000",
                          showErrorMessage=True, errorTitle="Nieprawidłowa cena",
                          error="Wpisz cenę opakowania w zł (liczba od 1 do 5000), bez symbolu waluty.",
                          showInputMessage=True, promptTitle="Cena", prompt="Cena brutto za opakowanie w zł. Zostaw puste, jeśli apteka nie prowadzi tej odmiany.")
dv_qty = DataValidation(type="list", formula1='"dużo (15+),średnio (5–14),mało (1–4),brak"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Stan", error="Wybierz z listy: dużo, średnio, mało albo brak.",
                        showInputMessage=True, promptTitle="Stan", prompt="dużo = 15 i więcej op. · średnio = 5–14 · mało = 1–4 · brak = 0. Zmieniaj tylko, gdy stan przejdzie do innej grupy.")
dv_pack = DataValidation(type="list", formula1='"5,10,15"', showErrorMessage=True,
                         errorTitle="Gramatura", error="Wybierz 5, 10 lub 15 g.")
dv_date = DataValidation(type="date", operator="greaterThan", formula1="45658",
                         showErrorMessage=True, errorTitle="Data", error="Wpisz datę w formacie DD.MM.RRRR.")
for dv in (dv_price, dv_qty, dv_pack, dv_date):
    ws.add_data_validation(dv)
dv_price.add(f"F{r0}:F{last}")
dv_qty.add(f"G{r0}:G{last}")
for i, r in enumerate(rows):
    if r["pack"] is None:
        dv_pack.add(f"E{r0 + i}")
dv_date.add("C4")
dv_cnt = DataValidation(type="whole", operator="between", formula1="0", formula2="999", allow_blank=True,
                        showErrorMessage=True, errorTitle="Liczba opakowań", error="Wpisz liczbę całkowitą albo zostaw puste.",
                        showInputMessage=True, promptTitle="Opcjonalnie", prompt="Dla chętnych: dokładna liczba opakowań. Wtedy nie trzeba wybierać stanu z listy.")
ws.add_data_validation(dv_cnt)
dv_cnt.add(f"H{r0}:H{last}")
dv_new = DataValidation(type="list", formula1='"TAK"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Nowość", error="Wybierz TAK albo zostaw puste.",
                        showInputMessage=True, promptTitle="Nowość", prompt="TAK = na wydruku pojawi się znacznik ★ NOWOŚĆ.")
ws.add_data_validation(dv_new)
dv_new.add(f"I{r0}:I{last}")

# formatowanie warunkowe kolumny Stan
rng = f"G{r0}:H{last}"
red = dict(fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(color="9C0006", bold=True))
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($F{r0}<>"",$G{r0}="",$H{r0}="")'], **red))
ws.conditional_formatting.add(rng, FormulaRule(formula=[
    f'AND($G{r0}<>"",ISNUMBER($H{r0}),$G{r0}<>IF($H{r0}=0,"brak",IF($H{r0}<=4,"mało (1–4)",IF($H{r0}<=14,"średnio (5–14)","dużo (15+)"))))'], **red))
gr = f"G{r0}:G{last}"
for pref, fg, fc in (("dużo", "C6EFCE", "006100"), ("średnio", "E2F0D9", "38761D"),
                     ("mało", "FCE4D6", "C55A11"), ("brak", "FFEB9C", "9C5700")):
    ws.conditional_formatting.add(gr, FormulaRule(formula=[f'LEFT($G{r0},{len(pref)})="{pref}"'],
                                  fill=PatternFill("solid", fgColor=fg), font=Font(color=fc, bold=True)))

ws.freeze_panes = ws.cell(row=r0, column=4)
ws.auto_filter.ref = f"A{HDR}:I{last}"
ws.protection.sheet = True
ws.protection.autoFilter = False
ws.protection.sort = False
ws.protection.formatColumns = False
ws.print_title_rows = f"{HDR}:{HDR}"
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# kolumna pomocnicza (ukryta): numer kolejny odmian z wpisaną ceną -> lista do druku
for n in range(r0, last + 1):
    c = ws.cell(row=n, column=10, value=f'=IF(F{n}="","",COUNT($F${r0}:F{n}))')
    c.font = Font(name=F, size=8, color="BFBFBF")
ws.column_dimensions["J"].hidden = True
ws.column_dimensions["K"].hidden = True

# ---------------- Arkusz: Do druku (prosty, surowy układ) ----------------
wp = wb.create_sheet("Do druku")
wp.sheet_view.showGridLines = False
SRC = "'Stany i ceny'"
pw = {"A": 22, "B": 32, "C": 8, "D": 11, "E": 14, "F": 11, "G": 16}
for k, w in pw.items():
    wp.column_dimensions[k].width = w
wp.merge_cells("A1:G1"); wp.merge_cells("A2:G2"); wp.merge_cells("A3:G3")
wp["A1"] = f'=IF({SRC}!C3="","Surowce farmaceutyczne konopi",{SRC}!C3)'
wp["A1"].font = Font(name=F, size=16, bold=True, color="1F4E78")
wp["A2"] = "Dostępne surowce farmaceutyczne konopi: ceny i dostępność"
wp["A2"].font = Font(name=F, size=11, color="404040")
wp["A3"] = f'=IF({SRC}!C4="","Stan na: —","Stan na: "&TEXT({SRC}!C4,"DD.MM.YYYY"))'
wp["A3"].font = Font(name=F, size=11, bold=True)
for c in ("A1", "A2", "A3"):
    wp[c].alignment = Alignment(horizontal="center")
PH = 5
ph = ["Producent", "Nazwa odmiany", "THC", "Opak. (g)", "Cena (zł/op.)", "zł/g", "Dostępność"]
for i, h in enumerate(ph, start=1):
    c = wp.cell(row=PH, column=i, value=h)
    c.font = Font(name=F, size=10, bold=True, color="FFFFFF"); c.fill = fill_head
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = border
wp.row_dimensions[PH].height = 24
rngI = f"{SRC}!$J${r0}:$J${last}"
def idx(col, k):
    return f"INDEX({SRC}!${col}${r0}:${col}${last},MATCH({k},{rngI},0))"
npr = len(rows)
for k in range(1, npr + 1):
    n = PH + k
    wp.cell(row=n, column=1, value=f'=IFERROR({idx("B",k)},"")')
    wp.cell(row=n, column=2, value=f'=IFERROR({idx("C",k)}&IF({idx("I",k)}="TAK","  ★ NOWOŚĆ",""),"")')
    wp.cell(row=n, column=3, value=f'=IFERROR({idx("D",k)},"")')
    wp.cell(row=n, column=4, value=f'=IFERROR({idx("E",k)},"")')
    wp.cell(row=n, column=5, value=f'=IFERROR({idx("F",k)},"")')
    wp.cell(row=n, column=6, value=f'=IFERROR(IF(OR(D{n}="",E{n}=""),"",E{n}/D{n}),"")')
    h_ = idx("K", k)
    wp.cell(row=n, column=7, value=f'=IFERROR(IF(OR(LEFT({h_},4)="dużo",LEFT({h_},7)="średnio"),"dostępna",IF(LEFT({h_},4)="mało","ostatnie sztuki",IF({h_}="brak","chwilowy brak","—"))),"")')
    for col in range(1, 8):
        c = wp.cell(row=n, column=col)
        c.font = Font(name=F, size=10, bold=(col == 2))
        c.alignment = Alignment(horizontal="left" if col <= 2 else "center", vertical="center")
    wp.cell(row=n, column=5).number_format = '#,##0.00 "zł"'
    wp.cell(row=n, column=6).number_format = '#,##0.00'
plast = PH + npr
wp.conditional_formatting.add(f"A{PH+1}:G{plast}", FormulaRule(formula=[f'AND($B{PH+1}<>"",MOD(ROW(),2)=0)'],
                              fill=PatternFill("solid", fgColor="F2F2F2"), border=border))
wp.conditional_formatting.add(f"A{PH+1}:G{plast}", FormulaRule(formula=[f'$B{PH+1}<>""'], border=border))
wp.print_title_rows = f"{PH}:{PH}"
wp.page_setup.orientation = "portrait"
wp.page_setup.paperSize = wp.PAPERSIZE_A4
wp.page_setup.fitToWidth = 1
wp.page_setup.fitToHeight = 0
wp.sheet_properties.pageSetUpPr.fitToPage = True
wp.print_options.horizontalCentered = True
wp.page_margins.left = wp.page_margins.right = 0.4
wp.page_margins.top = wp.page_margins.bottom = 0.5
wp.oddFooter.center.text = "Ceny i dostępność mogą się zmienić. Prosimy o potwierdzenie przy okienku."
wp.oddFooter.center.size = 8
from openpyxl.workbook.defined_name import DefinedName
wp.defined_names["_xlnm.Print_Area"] = DefinedName("_xlnm.Print_Area",
    attr_text=f"OFFSET('Do druku'!$A$1,0,0,{PH}+MAX(1,COUNT({rngI})),7)")
wp.protection.sheet = True

# ---------------- Arkusz 2: Inna odmiana ----------------
w2 = wb.create_sheet("Inna odmiana")
w2.sheet_view.showGridLines = False
w2["A1"] = "Odmiany spoza listy"
w2["A1"].font = Font(name=F, size=14, bold=True, color="1F4E78")
w2["A2"] = "Wpisz tu surowce, których nie ma w arkuszu „Stany i ceny”. Dodam je do listy."
w2["A2"].font = Font(name=F, size=10, italic=True, color="595959")
h2 = ["Producent", "Nazwa odmiany", "THC", "Opakowanie (g)", "Cena (zł/op.)", "Stan", "Kod EAN (opcjonalnie)", "Uwagi"]
w2w = [22, 30, 9, 14, 14, 15, 20, 30]
for i, (h, w) in enumerate(zip(h2, w2w), start=1):
    c = w2.cell(row=4, column=i, value=h)
    c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
    c.fill = fill_head
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border
    w2.column_dimensions[c.column_letter].width = w
w2.row_dimensions[4].height = 30
for n in range(5, 25):
    for col in range(1, 9):
        c = w2.cell(row=n, column=col)
        c.fill = fill_in
        c.border = border
        c.font = Font(name=F, size=10, color="0000FF")
    w2.cell(row=n, column=5).number_format = '#,##0.00 "zł"'
    w2.cell(row=n, column=7).number_format = "@"
dv2p = DataValidation(type="decimal", operator="between", formula1="1", formula2="5000", showErrorMessage=True,
                      error="Cena w zł, liczba od 1 do 5000.")
dv2q = DataValidation(type="list", formula1='"dużo (15+),średnio (5–14),mało (1–4),brak"', allow_blank=True, showErrorMessage=True,
                      error="Wybierz z listy: dużo, średnio, mało albo brak.")
dv2g = DataValidation(type="list", formula1='"5,10,15"', showErrorMessage=True, error="Wybierz 5, 10 lub 15 g.")
for dv in (dv2p, dv2q, dv2g):
    w2.add_data_validation(dv)
dv2p.add("E5:E24"); dv2q.add("F5:F24"); dv2g.add("D5:D24")
w2.freeze_panes = "A5"

# ---------------- Arkusz 3: Instrukcja ----------------
w3 = wb.create_sheet("Instrukcja")
w3.sheet_view.showGridLines = False
w3.column_dimensions["A"].width = 3
iw = [20, 22, 8, 13, 13, 14, 16]
for k, w in zip("BCDEFGH", iw):
    w3.column_dimensions[k].width = w
lines = [
    ("Instrukcja wypełniania", "title"),
    ("", None),
    ("Wypełniamy tylko żółte pola. Szare kolumny (ID, producent, nazwa, THC) są zablokowane.", "b"),
    ("", None),
    ("1. W arkuszu „Stany i ceny” wpisz nazwę apteki i datę stanu (pola na górze).", None),
    ("2. Przy każdej odmianie, którą apteka prowadzi, wpisz cenę za opakowanie (zł) i wybierz stan z listy. Dla chętnych: zamiast stanu możesz wpisać dokładną liczbę opakowań.", None),
    ("3. Odmiany, których apteka nie prowadzi, zostaw puste.", None),
    ("    Ta sama odmiana w różnych stężeniach THC lub gramaturach ma osobne wiersze: cenę i stan wpisz w wierszu, który się zgadza.", None),
    ("4. Przy kolejnej aktualizacji zmieniaj tylko to, co się zmieniło. Stan zmieniaj dopiero, gdy przejdzie do innej grupy (np. z „średnio” na „mało”).", None),
    ("5. Surowiec spoza listy wpisz w arkuszu „Inna odmiana”. Nowość w ofercie oznacz „TAK” w kolumnie Nowość.", None),
    ("6. Zapisz plik i wyślij go jako załącznik na: dostepnosc.cannabis@gmail.com", None),
    ("7. Chcesz cennik przy kasie? Otwórz arkusz „Do druku” i wydrukuj (Ctrl+P). Lista układa się sama z wpisanych cen.", None),
    ("8. Chcesz cennik ozdobić po swojemu (kolory, emotki, zdjęcia)? W arkuszu „Do druku” zaznacz tabelę od nagłówka do ostatniej odmiany, skopiuj (Ctrl+C) i wklej do Worda lub Dokumentów Google (Ctrl+V). Tam edytujesz dowolnie.", None),
    ("    Uwaga: kopia w Wordzie nie aktualizuje się sama. Po zmianie cen skopiuj tabelę ponownie.", None),
    ("", None),
    ("Stany:", "b"),
    ("•  Cena pusta  →  nie prowadzimy", None),
    ("•  dużo  →  15 i więcej opakowań", None),
    ("•  średnio  →  5–14 opakowań", None),
    ("•  mało  →  1–4 opakowania (na wydruku: „ostatnie sztuki”)", None),
    ("•  brak  →  odmiana w ofercie, chwilowo brak (czekamy na dostawę)", None),
    ("•  Cena wpisana, a stan i liczba puste  →  pola podświetlą się na czerwono", None),
    ("•  Wpisana liczba niezgodna z wybranym stanem (np. 8 i „mało”)  →  pola podświetlą się na czerwono", None),
    ("", None),
    ("Przykład poprawnie wypełnionego wiersza:", "b"),
]
for i, (t, kind) in enumerate(lines, start=1):
    w3.merge_cells(start_row=i, start_column=2, end_row=i, end_column=8)
    c = w3.cell(row=i, column=2, value=t)
    c.font = Font(name=F, size=14 if kind == "title" else 11, bold=kind in ("title", "b"),
                  color="1F4E78" if kind == "title" else "000000")
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if len(t) > 110:
        w3.row_dimensions[i].height = 30
hr = len(lines) + 1
ex_h = ["Producent", "Nazwa odmiany", "THC", "Opakowanie (g)", "Cena (zł/op.)", "Stan", "Liczba op. (opcjonalnie)"]
ex_v = ["Aurora", "Sourdough", "29%", 10, 360, "średnio (5–14)", None]
for j, (h, v) in enumerate(zip(ex_h, ex_v)):
    a = w3.cell(row=hr, column=2 + j, value=h)
    a.font = Font(name=F, size=10, bold=True, color="FFFFFF"); a.fill = fill_head
    a.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); a.border = border
    b = w3.cell(row=hr + 1, column=2 + j, value=v)
    b.border = border; b.alignment = Alignment(horizontal="center")
    inp = h in ("Cena (zł/op.)", "Stan")
    b.font = Font(name=F, size=10, color="0000FF" if inp else "000000", bold=(h == "Nazwa odmiany"))
    b.fill = fill_in if inp else fill_lock
    if h == "Cena (zł/op.)":
        b.number_format = '#,##0.00 "zł"'
    if h == "Stan":
        b.fill = PatternFill("solid", fgColor="E2F0D9"); b.font = Font(name=F, size=10, color="38761D", bold=True)
    if h.startswith("Liczba"):
        b.fill = PatternFill("solid", fgColor="FFFBEA")
w3.row_dimensions[hr].height = 30
foot = hr + 3
w3.merge_cells(start_row=foot, start_column=2, end_row=foot, end_column=8)
w3.cell(row=foot, column=2, value="Pytania: lek. Michał Guzik, dostepnosc.cannabis@gmail.com").font = Font(name=F, size=10, italic=True, color="595959")
w3.page_setup.fitToWidth = 1
w3.page_setup.fitToHeight = 0
w3.sheet_properties.pageSetUpPr.fitToPage = True
w3.protection.sheet = True

# ---------------- meta (ukryty) ----------------
w4 = wb.create_sheet("meta")
w4["A1"], w4["B1"] = "template_version", "9"
w4["A2"], w4["B2"] = "strains_json_last_updated", SRC_DATE
w4["A3"], w4["B3"] = "rows", len(rows)
w4["A4"], w4["B4"] = "data_first_row", r0
w4.sheet_state = "hidden"
for row in w4.iter_rows():
    for c in row:
        c.font = Font(name=F, size=10)

wb.move_sheet("Instrukcja", offset=-3)
wb.active = 1
wb.save(out)
print(len(rows), "rows; unknown-pack rows:", sum(1 for r in rows if r["pack"] is None))
